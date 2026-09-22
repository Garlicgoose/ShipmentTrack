# -*- coding: utf-8 -*-
"""FedEx webpage POD: user submits/searches; app only fills after a click and prints."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from urllib.parse import urlparse
from openpyxl import load_workbook

from modules.fedex_web_pod import (
    DETAIL_CONTENT_RE,
    _body_text,
    _main_page_ready,
    _print_current_page,
    is_valid_pdf,
    normalize_tracking_number,
    output_paths,
)
from modules.real_browser import RealBrowserController
from units import get_data_path


BLOCK_PAGE_RE = re.compile(
    r"too many requests|access denied|temporarily unavailable|service unavailable|"
    r"captcha|unusual traffic|请求过多|访问被拒绝|验证码|系统错误",
    re.I,
)

# Only a real pointer click on a tracking input arms typing. The Playwright
# worker sends actual key events afterwards; direct DOM value replacement may
# leave FedEx's React TRACK button disabled. Nothing clicks or submits TRACK.
ARM_INPUT_SCRIPT = r"""
(number) => {
  const key = '__shipmentTrackManualFill';
  if (window[key]) document.removeEventListener('pointerdown', window[key], true);
  window.__shipmentTrackManualPending = null;
  const handler = event => {
    const field = event.target?.closest?.('input, textarea');
    if (!field) return;
    const hint = [field.id, field.name, field.placeholder, field.ariaLabel]
      .filter(Boolean).join(' ').toLowerCase();
    if (!/track|tracking|运单|追踪/.test(hint)) return;
    setTimeout(() => {
      window.__shipmentTrackManualPending = number;
      document.removeEventListener('pointerdown', handler, true);
      window[key] = null;
    }, 0);
  };
  window[key] = handler;
  document.addEventListener('pointerdown', handler, true);
}
"""


@dataclass(frozen=True)
class ManualPodResult:
    number: str
    main_pdf: str
    detail_pdf: str


class ManualFedExPodSession:
    def __init__(self, playwright, browser_path, browser_type, output_dir, log=None):
        self.output_dir = Path(output_dir)
        kind = str(browser_type or "edge").casefold()
        self.controller = RealBrowserController(
            playwright,
            browser_path,
            get_data_path() / "browser_profiles" / kind / "fedex_manual",
            browser_name="Google Chrome" if kind == "chrome" else "Microsoft Edge",
            minimize=False,
            log_func=log,
        )
        self.context = None
        self.page = None
        self._armed_page = None
        self._armed_url = ""

    def start(self):
        self.context, self.page = self.controller.start()

    def _find_fedex_page(self):
        """Follow the page the user opened; never navigate on their behalf."""
        if self.context is None:
            return None
        try:
            pages = list(self.context.pages)
        except Exception:
            return None
        for page in reversed(pages):
            try:
                host = (urlparse(str(page.url)).hostname or "").casefold()
                if host == "fedex.com" or host.endswith(".fedex.com"):
                    self.page = page
                    return page
            except Exception:
                continue
        return None

    def prepare(self, number):
        number = normalize_tracking_number(number)
        if self.context is None:
            self.start()
        self._armed_page = None
        self._armed_url = ""
        return number

    def main_ready(self, number):
        page = self._find_fedex_page()
        if page is None:
            return False
        text = _body_text(page)
        if BLOCK_PAGE_RE.search(text):
            raise RuntimeError("FedEx 页面出现限流、验证码或服务错误；请人工处理后重试")
        try:
            return _main_page_ready(page, number)
        except Exception:
            # User-initiated navigation can invalidate the current JS context.
            return False

    def fill_after_user_click(self, number):
        """Type only after the user actually clicks FedEx's tracking field."""
        page = self._find_fedex_page()
        if page is None:
            return False
        try:
            url = str(page.url)
            if page is not self._armed_page or url != self._armed_url:
                page.evaluate(ARM_INPUT_SCRIPT, number)
                self._armed_page, self._armed_url = page, url
            pending = page.evaluate(
                """() => {
                    const number = window.__shipmentTrackManualPending;
                    window.__shipmentTrackManualPending = null;
                    return number || '';
                }"""
            )
        except Exception:
            self._armed_page = None
            return False
        if pending != number:
            return False
        try:
            page.keyboard.press("Control+A")
            page.keyboard.type(number, delay=60)
        except Exception:
            # A user may submit/navigate while key events are in flight.
            self._armed_page = None
            return False
        return True

    def save_main(self, number):
        if not self.main_ready(number):
            raise RuntimeError("尚未看到当前运单的查询结果，不保存主页")
        main_pdf, _ = output_paths(self.output_dir, number)
        snapshot = (_body_text(self.page), str(self.page.url))
        _print_current_page(self.context, self.page, main_pdf)
        return str(main_pdf), snapshot

    def detail_ready(self, number, snapshot):
        page = self._find_fedex_page()
        if page is None:
            return False
        text = _body_text(page)
        if BLOCK_PAGE_RE.search(text):
            raise RuntimeError("FedEx 详情页出现限流、验证码或服务错误")
        before_text, before_url = snapshot
        compact = re.sub(r"[\s-]+", "", text)
        return (
            number in compact
            and text != before_text
            and bool(DETAIL_CONTENT_RE.search(text))
            and (str(page.url) != before_url or abs(len(text) - len(before_text)) >= 80)
        )

    def save_detail(self, number, snapshot):
        if not self.detail_ready(number, snapshot):
            raise RuntimeError("尚未看到当前运单的详情内容，不保存详情页")
        _, detail_pdf = output_paths(self.output_dir, number)
        _print_current_page(self.context, self.page, detail_pdf)
        if not is_valid_pdf(detail_pdf):
            raise RuntimeError("FedEx 详情 PDF 无效")
        return str(detail_pdf)

    def close(self):
        self.controller.close()
        self.context = self.page = None
        self._armed_page = None


def update_tracking_result_file(result_file, pod_result: ManualPodResult) -> bool:
    """Write two manually saved POD paths back to the existing tracking result."""
    path = Path(result_file)
    if not path.is_file():
        return False
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        headers = {str(cell.value or "").strip(): cell.column for cell in sheet[1]}
        required = {"运单号", "快递公司", "POD文件", "POD详情文件"}
        if not required.issubset(headers):
            return False
        for row in range(2, sheet.max_row + 1):
            number = str(sheet.cell(row, headers["运单号"]).value or "").strip()
            carrier = str(sheet.cell(row, headers["快递公司"]).value or "").strip()
            if number == pod_result.number and carrier.casefold() == "fedex":
                sheet.cell(row, headers["POD文件"], pod_result.main_pdf)
                sheet.cell(row, headers["POD详情文件"], pod_result.detail_pdf)
                workbook.save(path)
                return True
        return False
    finally:
        workbook.close()


def refresh_pod_audit(result_file, audit_file, sample_rates=None):
    """Re-sample all available PODs after manual FedEx pages are saved."""
    from modules.pod_audit import audit_pod_sample
    from modules.tracking_runner import _aggregate_audit_results

    path = Path(result_file)
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        headers = {str(cell.value or "").strip(): cell.column for cell in sheet[1]}
        required = {"运单号", "快递公司", "状态", "POD文件", "POD详情文件", "POD抽查"}
        if not required.issubset(headers):
            raise ValueError("跟踪结果缺少 POD 抽查需要的列")
        rows = []
        for index in range(2, sheet.max_row + 1):
            item = {key: sheet.cell(index, column).value or "" for key, column in headers.items()}
            item["is_delivered"] = bool(item.get("POD文件"))
            rows.append(item)
        audits = audit_pod_sample(rows, audit_file, sample_rates=sample_rates)
        outcomes = _aggregate_audit_results(audits)
        for index in range(2, sheet.max_row + 1):
            number = str(sheet.cell(index, headers["运单号"]).value or "")
            sheet.cell(index, headers["POD抽查"], outcomes.get(number, ""))
        workbook.save(path)
        return len(audits)
    finally:
        workbook.close()
