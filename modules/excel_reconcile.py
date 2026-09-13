# -*- coding: utf-8 -*-
"""检验表与 Droplist 的合并、分类和按日核对。"""
from __future__ import annotations

from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from pathlib import Path
import re
from typing import Callable, Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from modules.settings_store import FilenameMapper, FilenameMappingRule


ProgressCallback = Optional[Callable[[int], None]]


@dataclass(frozen=True)
class ReconcileRow:
    date: str
    target_type: str
    inspect_quantity: float
    droplist_quantity: float
    difference: float
    result: str


@dataclass(frozen=True)
class ExcelReconcileResult:
    inspect_output_file: Path
    droplist_output_file: Path
    inspect_files: int
    droplist_files: int
    inspect_rows: int
    droplist_rows: int
    rows: tuple[ReconcileRow, ...]
    issues: tuple[tuple[str, str, str], ...]


def extract_date_from_name(name: str) -> tuple[tuple[int, int, int], str]:
    """从文件名中提取 YYYY.M.D 或 M.D，返回排序键与显示值。"""
    filename = Path(str(name or "")).name
    stem = re.sub(r"\.(?:xlsx|xlsm|xls)$", "", filename, flags=re.IGNORECASE)
    match = re.search(
        r"(?<!\d)(?:(20\d{2})[._-])?(\d{1,2})[._-](\d{1,2})(?!\d)",
        stem,
    )
    if not match:
        return (0, 0, 0), ""
    year = int(match.group(1) or 0)
    month = int(match.group(2))
    day = int(match.group(3))
    if not (1 <= month <= 12 and 1 <= day <= 31):
        return (0, 0, 0), ""
    label = f"{year}.{month}.{day}" if year else f"{month}.{day}"
    return (year, month, day), label


def _copy_style(source, target) -> None:
    if not source.has_style:
        return
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)
    if getattr(source, "comment", None) is not None:
        target.comment = copy(source.comment)
    if getattr(source, "hyperlink", None) is not None:
        target._hyperlink = copy(source.hyperlink)


def _safe_value(sheet, row: int, column: int):
    cell = sheet.cell(row=row, column=column)
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged and (
            row != merged.min_row or column != merged.min_col
        ):
            return None
    return cell.value


def _quantity(value) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _is_empty_row(sheet, row: int, max_column: int) -> bool:
    return all(
        _safe_value(sheet, row, column) in (None, "")
        for column in range(1, max_column + 1)
    )


def _last_non_empty_row(sheet, start_row: int, max_column: int) -> int:
    """返回从 start_row 起最后一个有内容的行号；没有内容时返回 start_row - 1。"""
    for row in range(sheet.max_row, start_row - 1, -1):
        if not _is_empty_row(sheet, row, max_column):
            return row
    return start_row - 1


def _is_overseas_truck_filename(name: str) -> bool:
    """识别“国外第1车 / 国外出货第一车”等检验表文件名。"""
    stem = Path(str(name or "")).stem
    return bool(
        re.search(
            r"国外\s*(?:出货\s*)?第\s*[0-9一二三四五六七八九十百零〇两]+\s*车",
            stem,
            flags=re.IGNORECASE,
        )
    )


def _is_droplist_end_row(sheet, row: int, max_column: int) -> bool:
    first = _safe_value(sheet, row, 1)
    if isinstance(first, str):
        folded = first.casefold()
        if "by fed-ex" in folded or "total" in folded:
            return True
    return _is_empty_row(sheet, row, max_column)


def _is_droplist_data_sheet(sheet) -> bool:
    """有效明细页不依赖页签名，只认第 3 行的业务表头。"""
    if sheet.max_row < 4:
        return False
    headers = {
        str(_safe_value(sheet, 3, column) or "").strip().casefold()
        for column in range(1, sheet.max_column + 1)
    }
    return "s/o" in headers and "qty" in headers


def _copy_column_layout(source_sheet, target_sheet, max_column):
    for column in range(1, max_column + 1):
        letter = get_column_letter(column)
        source = source_sheet.column_dimensions[letter]
        target = target_sheet.column_dimensions[letter]
        if source.width is not None:
            target.width = max(target.width or 0, source.width)
        target.hidden = target.hidden or source.hidden
        target.bestFit = target.bestFit or source.bestFit
        target.outlineLevel = max(target.outlineLevel, source.outlineLevel)


def _append_source_row(source_sheet, source_row, target_sheet, target_row, max_column):
    source_dimension = source_sheet.row_dimensions[source_row]
    target_dimension = target_sheet.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed
    for column in range(1, max_column + 1):
        source = source_sheet.cell(row=source_row, column=column)
        value = _safe_value(source_sheet, source_row, column)
        if source.data_type == "f" and isinstance(value, str):
            try:
                value = Translator(
                    value,
                    origin=source.coordinate,
                ).translate_formula(get_column_letter(column) + str(target_row))
            except (TypeError, ValueError):
                pass
        target = target_sheet.cell(
            row=target_row,
            column=column,
            value=value,
        )
        _copy_style(source, target)


def _copy_merged_ranges(source_sheet, target_sheet, row_map, max_column):
    """把完整落在已复制行中的合并区域平移到目标表。"""
    for merged in source_sheet.merged_cells.ranges:
        if merged.max_col > max_column:
            continue
        source_rows = list(range(merged.min_row, merged.max_row + 1))
        if not all(row in row_map for row in source_rows):
            continue
        target_rows = [row_map[row] for row in source_rows]
        if target_rows != list(range(target_rows[0], target_rows[0] + len(target_rows))):
            continue
        target_sheet.merge_cells(
            start_row=target_rows[0],
            start_column=merged.min_col,
            end_row=target_rows[-1],
            end_column=merged.max_col,
        )


def _append_metadata_headers(sheet, row, start_column):
    for offset, value in enumerate(
        ("类型", "归总类别", "日期", "来源文件", "映射备注")
    ):
        sheet.cell(row=row, column=start_column + offset, value=value)


def _append_metadata(
    sheet,
    row,
    start_column,
    display_type,
    target_type,
    date,
    source,
    note,
):
    values = (display_type, target_type, date, source, note)
    for offset, value in enumerate(values):
        sheet.cell(row=row, column=start_column + offset, value=value)


def _resolve_date(file: Path) -> tuple[tuple[int, int, int], str]:
    key, label = extract_date_from_name(file.stem)
    if label:
        return key, label
    return extract_date_from_name(file.parent.name)


def _xlsx_files(folder: Path, recursive: bool, excluded: set[Path]) -> list[Path]:
    iterator = folder.rglob("*.xlsx") if recursive else folder.glob("*.xlsx")
    files = []
    for file in iterator:
        if file.name.startswith("~$"):
            continue
        try:
            resolved = file.resolve()
        except OSError:
            resolved = file.absolute()
        if resolved not in excluded:
            files.append(file)
    return files


def _merge_inspect(
    folder: Path,
    output_sheet,
    mapper: FilenameMapper,
    excluded: set[Path],
    issues: list[tuple[str, str, str]],
) -> tuple[
    dict[tuple[str, str], float],
    dict[tuple[str, str, str], float],
    int,
    int,
]:
    files = _xlsx_files(folder, recursive=False, excluded=excluded)
    files.sort(key=lambda file: (_resolve_date(file)[0], file.name.casefold()))
    if not files:
        raise FileNotFoundError("检验表文件夹中没有可处理的 .xlsx 文件")

    totals: dict[tuple[str, str], float] = defaultdict(float)
    display_totals: dict[tuple[str, str, str], float] = defaultdict(float)
    target_row = 1
    fixed_columns = 0
    data_rows = 0
    processed_files = 0

    for file_index, file in enumerate(files):
        workbook = load_workbook(file, data_only=False)
        source = workbook.active
        max_column = source.max_column
        _copy_column_layout(source, output_sheet, max_column)
        if file_index == 0:
            output_sheet.sheet_view.showGridLines = source.sheet_view.showGridLines
        last_source_row = _last_non_empty_row(source, 2, max_column)
        if last_source_row < 2:
            issues.append(("检验表", file.name, "没有数据行"))
            workbook.close()
            continue
        processed_files += 1
        row_map = {}
        if not fixed_columns:
            fixed_columns = max_column
            row_map[1] = target_row
            _append_source_row(source, 1, output_sheet, target_row, fixed_columns)
            _append_metadata_headers(output_sheet, target_row, fixed_columns + 1)
            target_row += 1
        elif max_column != fixed_columns:
            issues.append(
                ("检验表", file.name, f"列数 {max_column} 与首个文件 {fixed_columns} 不一致")
            )

        match = mapper.match(file.name)
        overseas_truck = _is_overseas_truck_filename(file.name)
        if match.target_type in {"光联", "MPO"}:
            comparison_type = match.target_type
        elif overseas_truck:
            comparison_type = "光联"
        else:
            comparison_type = "未识别"

        display_type = str(match.display_type or "").strip()
        mapping_note = match.note
        if overseas_truck and (not match.matched or display_type in {"", "未识别"}):
            # “国外第…车”没有配置类型时，默认属于光联。
            display_type = "光联"
        if overseas_truck and not match.matched:
            mapping_note = "默认规则：国外第…车归类为光联"

        _, date_label = _resolve_date(file)
        if not match.matched and not overseas_truck:
            issues.append(("检验表", file.name, match.note))
        if not date_label:
            issues.append(("检验表", file.name, "文件名和父文件夹均无法识别日期"))

        # 只复制到最后一个非空行，去掉每个检验表末尾多余的空行；
        # 文件内部原本存在的空行仍然保留。
        for source_row in range(2, last_source_row + 1):
            row_map[source_row] = target_row
            empty_row = _is_empty_row(source, source_row, max_column)
            _append_source_row(
                source,
                source_row,
                output_sheet,
                target_row,
                min(max_column, fixed_columns),
            )
            if not empty_row:
                _append_metadata(
                    output_sheet,
                    target_row,
                    fixed_columns + 1,
                    display_type,
                    comparison_type,
                    date_label,
                    file.name,
                    mapping_note,
                )
                quantity = _quantity(_safe_value(source, source_row, 6))
                totals[(date_label, comparison_type)] += quantity
                display_totals[(date_label, display_type, comparison_type)] += quantity
            target_row += 1
            if not empty_row:
                data_rows += 1
        _copy_merged_ranges(source, output_sheet, row_map, fixed_columns)
        workbook.close()

    return totals, display_totals, processed_files, data_rows


def _merge_droplist(
    folder: Path,
    output_sheet,
    mapper: FilenameMapper,
    excluded: set[Path],
    issues: list[tuple[str, str, str]],
) -> tuple[dict[tuple[str, str], float], int, int]:
    candidates = _xlsx_files(folder, recursive=True, excluded=excluded)
    files = [
        file for file in candidates
        if file.stem.casefold().startswith("drop shipment list")
    ]
    files.sort(key=lambda file: (_resolve_date(file)[0], file.name.casefold()))
    if not files:
        raise FileNotFoundError("Droplist 文件夹中没有有效的 Drop shipment list 文件")

    totals: dict[tuple[str, str], float] = defaultdict(float)
    target_row = 1
    fixed_columns = 0
    data_rows = 0
    processed_files = 0

    for file in files:
        workbook = load_workbook(file, data_only=False)
        match = mapper.match(file.name)
        comparison_type = (
            match.target_type if match.target_type in {"光联", "MPO"} else "未识别"
        )
        _, date_label = _resolve_date(file)
        if not match.matched:
            issues.append(("Droplist", file.name, match.note))
        if not date_label:
            issues.append(("Droplist", file.name, "文件名和父文件夹均无法识别日期"))

        file_rows = 0
        for sheet_name in workbook.sheetnames[1:]:
            normalized_name = sheet_name.strip().casefold()
            if normalized_name == "address" or re.fullmatch(r"sheet\s*\d*", normalized_name):
                continue
            source = workbook[sheet_name]
            if not _is_droplist_data_sheet(source):
                continue
            max_column = source.max_column
            _copy_column_layout(source, output_sheet, max_column)
            row_map = {}
            if not fixed_columns:
                fixed_columns = max_column
                row_map[3] = target_row
                _append_source_row(source, 3, output_sheet, target_row, fixed_columns)
                _append_metadata_headers(output_sheet, target_row, fixed_columns + 1)
                target_row += 1
            elif max_column != fixed_columns:
                issues.append(
                    ("Droplist", file.name, f"列数 {max_column} 与首个数据表 {fixed_columns} 不一致")
                )

            for source_row in range(4, source.max_row + 1):
                if _is_droplist_end_row(source, source_row, max_column):
                    break
                row_map[source_row] = target_row
                _append_source_row(
                    source,
                    source_row,
                    output_sheet,
                    target_row,
                    min(max_column, fixed_columns),
                )
                _append_metadata(
                    output_sheet,
                    target_row,
                    fixed_columns + 1,
                    comparison_type,
                    comparison_type,
                    date_label,
                    file.name,
                    match.note,
                )
                totals[(date_label, comparison_type)] += _quantity(
                    _safe_value(source, source_row, 6)
                )
                target_row += 1
                data_rows += 1
                file_rows += 1
            _copy_merged_ranges(source, output_sheet, row_map, fixed_columns)
        if not file_rows:
            issues.append(("Droplist", file.name, "没有有效数据行"))
        else:
            processed_files += 1
        workbook.close()

    return totals, processed_files, data_rows


def _date_sort_key(label: str):
    return extract_date_from_name(label)[0]


def _build_reconcile_rows(inspect_totals, droplist_totals) -> tuple[ReconcileRow, ...]:
    rows = []
    keys = sorted(
        set(inspect_totals) | set(droplist_totals),
        key=lambda item: (_date_sort_key(item[0]), item[1]),
    )
    for date, target_type in keys:
        inspect = inspect_totals.get((date, target_type), 0.0)
        droplist = droplist_totals.get((date, target_type), 0.0)
        difference = inspect - droplist
        if (date, target_type) not in inspect_totals:
            result = "检验表缺少数据"
        elif (date, target_type) not in droplist_totals:
            result = "Droplist 缺少数据"
        elif abs(difference) < 1e-9:
            result = "一致"
        else:
            result = "不一致"
        rows.append(ReconcileRow(date, target_type, inspect, droplist, difference, result))
    return tuple(rows)


def _style_output(workbook) -> None:
    dark_fill = PatternFill("solid", fgColor="173F5F")
    light_fill = PatternFill("solid", fgColor="EAF2F7")
    # 合并明细页保留源文件格式；只格式化本程序新建的汇总和异常页。
    for sheet in (
        workbook["类型箱数"],
        workbook["核对汇总"],
        workbook["异常文件"],
    ):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2" if sheet.max_row > 1 else None
        for cell in sheet[1]:
            cell.fill = dark_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.auto_filter.ref = sheet.dimensions if sheet.max_row > 1 else None
        for column_cells in sheet.columns:
            values = [str(cell.value or "") for cell in column_cells[:200]]
            width = min(max(max((len(value) for value in values), default=8) + 2, 10), 36)
            sheet.column_dimensions[column_cells[0].column_letter].width = width

    summary = workbook["核对汇总"]
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 3).number_format = "#,##0.##"
        summary.cell(row, 4).number_format = "#,##0.##"
        summary.cell(row, 5).number_format = "#,##0.##"
        if summary.cell(row, 6).value == "一致":
            summary.cell(row, 6).fill = PatternFill("solid", fgColor="E8F6EE")
        else:
            summary.cell(row, 6).fill = PatternFill("solid", fgColor="FFF0E0")
    if summary.max_row > 1:
        for cell in summary[summary.max_row]:
            if cell.row > 1 and cell.value == "总计":
                cell.fill = light_fill


def merge_and_reconcile_excel(
    inspect_folder: Path,
    droplist_folder: Path,
    output_dir: Path,
    mapping_rules: Iterable[FilenameMappingRule],
    progress_callback: ProgressCallback = None,
) -> ExcelReconcileResult:
    inspect_folder = Path(inspect_folder)
    droplist_folder = Path(droplist_folder)
    output_dir = Path(output_dir)
    if not inspect_folder.is_dir():
        raise NotADirectoryError(f"检验表文件夹不存在：{inspect_folder}")
    if not droplist_folder.is_dir():
        raise NotADirectoryError(f"Droplist 文件夹不存在：{droplist_folder}")

    output_dir.mkdir(parents=True, exist_ok=True)
    inspect_output_file = output_dir / "合并检验表.xlsx"
    droplist_output_file = output_dir / "合并Droplist.xlsx"
    excluded = {inspect_output_file.resolve(), droplist_output_file.resolve()}
    mapper = FilenameMapper(mapping_rules)
    issues: list[tuple[str, str, str]] = []

    inspect_workbook = Workbook()
    inspect_sheet = inspect_workbook.active
    inspect_sheet.title = "合并检验表"
    droplist_workbook = Workbook()
    droplist_sheet = droplist_workbook.active
    droplist_sheet.title = "合并Droplist"

    inspect_totals, display_totals, inspect_files, inspect_rows = _merge_inspect(
        inspect_folder, inspect_sheet, mapper, excluded, issues
    )
    if progress_callback:
        progress_callback(45)
    droplist_totals, droplist_files, droplist_rows = _merge_droplist(
        droplist_folder, droplist_sheet, mapper, excluded, issues
    )
    if progress_callback:
        progress_callback(80)

    rows = _build_reconcile_rows(inspect_totals, droplist_totals)
    type_sheet = inspect_workbook.create_sheet("类型箱数")
    type_sheet.append(("日期", "类型", "箱数", "归总类别"))
    for (date, display_type, target_type), quantity in sorted(
        display_totals.items(),
        key=lambda item: (_date_sort_key(item[0][0]), item[0][1], item[0][2]),
    ):
        type_sheet.append((date, display_type, quantity, target_type))

    summary_sheet = inspect_workbook.create_sheet("核对汇总")
    summary_sheet.append(("日期", "类型", "检验表数量", "Droplist数量", "差异", "结果"))
    for row in rows:
        summary_sheet.append(
            (
                row.date,
                row.target_type,
                row.inspect_quantity,
                row.droplist_quantity,
                row.difference,
                row.result,
            )
        )

    issue_sheet = inspect_workbook.create_sheet("异常文件")
    issue_sheet.append(("来源", "文件", "问题"))
    for issue in issues:
        issue_sheet.append(issue)

    _style_output(inspect_workbook)
    inspect_workbook.save(inspect_output_file)
    droplist_workbook.save(droplist_output_file)
    if progress_callback:
        progress_callback(100)

    return ExcelReconcileResult(
        inspect_output_file=inspect_output_file,
        droplist_output_file=droplist_output_file,
        inspect_files=inspect_files,
        droplist_files=droplist_files,
        inspect_rows=inspect_rows,
        droplist_rows=droplist_rows,
        rows=rows,
        issues=tuple(issues),
    )
