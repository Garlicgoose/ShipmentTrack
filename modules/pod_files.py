# -*- coding: utf-8 -*-
"""Move POD PDFs first, then optionally replace colliding PDFs as a batch."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import os
import shutil
import tempfile
from typing import Callable, Iterable

from openpyxl import load_workbook


@dataclass(frozen=True)
class PodMoveResult:
    moved: tuple[tuple[str, str], ...]
    skipped: tuple[str, ...]
    errors: tuple[str, ...]


def _replace_pdf(source: Path, target: Path):
    # Stage a complete copy on the destination volume before atomically
    # replacing the old PDF. A failed copy leaves both original files intact.
    descriptor, temporary = tempfile.mkstemp(prefix=".pod-", suffix=".tmp", dir=target.parent)
    os.close(descriptor)
    staged = Path(temporary)
    try:
        shutil.copy2(source, staged)
        os.replace(staged, target)
        source.unlink()
    finally:
        if staged.exists():
            staged.unlink()


def move_pod_files(
    entries: Iterable[tuple[str, str]], destination,
    confirm_overwrite: Callable[[tuple[tuple[str, str], ...]], bool] | None = None,
) -> PodMoveResult:
    root = Path(destination).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    moved = []
    skipped = []
    errors = []
    seen = set()
    conflicts = []
    for _carrier, path_text in entries:
        source = Path(str(path_text or "")).expanduser()
        try:
            source = source.resolve()
        except OSError:
            source = source.absolute()
        key = str(source).casefold()
        if key in seen:
            continue
        seen.add(key)
        # Defense in depth: a POD entry must resolve to an existing PDF file.
        # Never pass directories (including directories named *.pdf) or other
        # source documents to shutil.move.
        if not source.is_file() or source.suffix.casefold() != ".pdf":
            skipped.append(str(source))
            continue
        if root == source.parent:
            skipped.append(str(source))
            continue
        target = root / source.name
        if target.exists():
            if target.is_file() and not target.is_symlink():
                conflicts.append((source, target))
            else:
                errors.append(f"{target}：目标不是普通 PDF 文件，无法覆盖")
            continue
        try:
            shutil.move(str(source), str(target))
            moved.append((str(source), str(target)))
        except OSError as exc:
            errors.append(f"{source}：{exc}")
    if conflicts:
        approved = bool(confirm_overwrite and confirm_overwrite(
            tuple((str(source), str(target)) for source, target in conflicts)
        ))
        for source, target in conflicts:
            if not approved:
                skipped.append(str(source))
                continue
            try:
                if not source.is_file() or target.is_symlink() or (target.exists() and not target.is_file()):
                    raise OSError("来源或目标文件已变化，无法覆盖")
                _replace_pdf(source, target)
                moved.append((str(source), str(target)))
            except OSError as exc:
                errors.append(f"{source}：{exc}")
    return PodMoveResult(tuple(moved), tuple(skipped), tuple(errors))


def update_tracking_pod_paths(result_file, moved_pairs) -> int:
    path = Path(result_file)
    if not path.is_file():
        return 0
    mapping = {str(old).casefold(): str(new) for old, new in moved_pairs}
    if not mapping:
        return 0
    workbook = load_workbook(path)
    updates = 0
    try:
        sheet = workbook.active
        headers = {str(cell.value or "").strip(): cell.column for cell in sheet[1]}
        columns = [headers[name] for name in ("POD文件", "POD详情文件") if name in headers]
        for row in range(2, sheet.max_row + 1):
            for column in columns:
                current = str(sheet.cell(row, column).value or "")
                replacement = mapping.get(current.casefold())
                if replacement:
                    sheet.cell(row, column, replacement)
                    updates += 1
        if updates:
            workbook.save(path)
        return updates
    finally:
        workbook.close()
