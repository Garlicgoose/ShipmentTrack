# -*- coding: utf-8 -*-
"""快递批量查询运行引擎：读 Excel -> 逐单查询 -> 输出结果 Excel。

日志格式（输出面板）：[1/100] DHL 4561879 Delivered
输出 Excel 列：运单号 / 快递公司 / 状态 / 抵达时间 / 用时(秒) / 备注
"""
import time
import traceback
from pathlib import Path

import pandas as pd

from modules.tracking_utils import (
    TRACKING_CARRIER_CONFIG,
    TRACKING_OUTPUT_COLUMNS,
    prepare_tracking_input_dataframe,
    normalize_tracking_arrival_date,
    TrackingCarrierSession,
)
from modules import fedex_module


def _is_delivered_truthy(value):
    """兼容各模块的返回值：DHL "Yes"/"No"，FedEx "Y"/""，EI/DSV/UPS True/False。"""
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in ("yes", "y", "true", "1")


def run_tracking(
    input_file,
    output_dir,
    ei_login_enabled=False,
    ei_email="",
    ei_password="",
    fedex_api_key="",
    fedex_api_secret="",
    chrome_path="",
    minimize_browser=True,
    save_pdf=True,
    log=None,
    progress=None,
):
    """执行批量查询。log(msg) / progress(value 0-100) 为回调（主线程安全）。"""
    log = log or (lambda msg: None)
    progress = progress or (lambda v: None)

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    pdf_root = output_path / "pdf"
    pdf_root.mkdir(parents=True, exist_ok=True)

    # 让各模块使用输出目录下的 PDF 目录
    for cfg in TRACKING_CARRIER_CONFIG.values():
        if cfg.get("api_based"):
            continue
        try:
            module = __import__(cfg["module"], fromlist=["x"])
            if hasattr(module, "PDF_DIR"):
                pdf_dir = pdf_root / cfg["pdf_subdir"]
                pdf_dir.mkdir(parents=True, exist_ok=True)
                setattr(module, "PDF_DIR", str(pdf_dir))
        except Exception:
            pass
    fedex_module.PDF_DIR = str(pdf_root / "FedEx")

    log("读取并清洗 Excel")
    work_df = prepare_tracking_input_dataframe(input_file)

    if work_df.empty:
        raise ValueError("清洗后没有可查询的运单号。")

    cleaned_file = output_path / "tracking_list_cleaned_sorted.xlsx"
    work_df.to_excel(cleaned_file, index=False, engine="openpyxl")
    log(f"已生成清洗排序文件：{cleaned_file}")

    results = []
    current_carrier = None
    session = None
    playwright = None

    try:
        from playwright.sync_api import sync_playwright

        total = len(work_df)

        for idx, row in work_df.iterrows():
            carrier = row["快递公司"]
            tracking_number = row["运单号"]

            # 换承运商时切换会话
            if carrier != current_carrier:
                if session is not None:
                    session.close()
                    session = None
                current_carrier = carrier
                if not TRACKING_CARRIER_CONFIG[carrier].get("api_based"):
                    if playwright is None:
                        playwright = sync_playwright().start()
                    session = TrackingCarrierSession(
                        playwright=playwright,
                        carrier=carrier,
                        output_dir=output_path,
                        ei_login_enabled=ei_login_enabled,
                        ei_email=ei_email,
                        ei_password=ei_password,
                        chrome_path=chrome_path,
                        minimize_browser=minimize_browser,
                        log_func=log,
                        save_pdf=save_pdf,
                    )
                    session.start()

            start_time = time.perf_counter()

            try:
                if TRACKING_CARRIER_CONFIG[carrier].get("api_based"):
                    # FedEx 官方 API，不走浏览器
                    raw_result = fedex_module.query_fedex_one(
                        tracking_number,
                        api_key=fedex_api_key,
                        api_secret=fedex_api_secret,
                        save_pdf=save_pdf,
                        pdf_dir=str(pdf_root / "FedEx"),
                    )
                else:
                    raw_result = session.query_one(tracking_number)

                elapsed_seconds = round(time.perf_counter() - start_time, 2)

                status = raw_result.get("status", "") or ""
                error = raw_result.get("error", "") or ""
                flag = raw_result.get("flag", "") or ""

                delivered = _is_delivered_truthy(raw_result.get("is_delivered"))
                arrival_time = ""
                if delivered:
                    arrival_time = normalize_tracking_arrival_date(
                        raw_result.get("arrival_time", "") or ""
                    )

                remark = " | ".join(x for x in (error, flag) if x)

                results.append({
                    "运单号": tracking_number,
                    "快递公司": carrier,
                    "状态": status,
                    "抵达时间": arrival_time,
                    "用时(秒)": elapsed_seconds,
                    "备注": remark,
                })

                log(f"[{idx + 1}/{total}] {carrier} {tracking_number} {status}")

            except Exception as e:
                elapsed_seconds = round(time.perf_counter() - start_time, 2)
                results.append({
                    "运单号": tracking_number,
                    "快递公司": carrier,
                    "状态": "Error",
                    "抵达时间": "",
                    "用时(秒)": elapsed_seconds,
                    "备注": str(e),
                })
                log(f"[{idx + 1}/{total}] {carrier} {tracking_number} Error: {e}")

            progress(int(((idx + 1) / total) * 100))
            time.sleep(0.3)

        if session is not None:
            session.close()
            session = None
    finally:
        if playwright is not None:
            try:
                playwright.stop()
            except Exception:
                pass

    result_df = pd.DataFrame(results, columns=TRACKING_OUTPUT_COLUMNS)

    output_file = output_path / "tracking_result.xlsx"
    result_df.to_excel(output_file, index=False, engine="openpyxl")

    # 运单号按文本储存，避免科学计数法
    try:
        from openpyxl import load_workbook
        wb = load_workbook(output_file)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
            for cell in row:
                cell.number_format = "@"
        wb.save(output_file)
    except Exception:
        pass

    log("=" * 60)
    log("全部完成")
    log(f"输出结果：{output_file}")
    log("=" * 60)

    return output_file
