# -*- coding: utf-8 -*-
"""按承运商配置比例抽查 POD，并输出统一的可追溯结果。"""
from __future__ import annotations

from dataclasses import dataclass
import math
from pathlib import Path
import random
import re
from typing import Callable, Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader
from modules.settings_store import DEFAULT_POD_AUDIT_RATES


DELIVERED_TERMS = (
    "delivered",
    "completed",
    "proof of delivery",
    "delivery confirmation",
)


@dataclass(frozen=True)
class PodAuditItem:
    tracking_number: str
    carrier: str
    pdf_file: str
    sample_reason: str
    tracking_status: str
    valid_pdf: bool
    tracking_found: bool
    delivered_found: bool
    status_field: str
    result: str
    details: str
    pod_kind: str = "POD"
    sample_rate: float = 0.0


def _normalize_search_text(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def inspect_pod(pdf_file, tracking_number, carrier="") -> PodAuditItem:
    path_text = str(pdf_file or "").strip()
    if not path_text:
        return PodAuditItem(
            tracking_number=str(tracking_number),
            carrier=str(carrier),
            pdf_file="",
            sample_reason="",
            tracking_status="",
            valid_pdf=False,
            tracking_found=False,
            delivered_found=False,
            status_field="",
            result="失败",
            details="POD文件路径为空",
        )
    path = Path(path_text)
    valid_pdf = False
    tracking_found = False
    delivered_found = False
    status_field = ""
    details = ""
    try:
        reader = PdfReader(str(path))
        valid_pdf = len(reader.pages) > 0
        page_text = []
        for page in reader.pages[:5]:
            page_text.append(page.extract_text() or "")
        text = "\n".join(page_text)
        folded = text.casefold()
        normalized = _normalize_search_text(text)
        candidates = {
            _normalize_search_text(tracking_number),
            _normalize_search_text(path.stem),
        }
        candidates.discard("")
        tracking_found = any(candidate in normalized for candidate in candidates)
        for line in text.splitlines():
            if any(term in line.casefold() for term in DELIVERED_TERMS):
                status_field = " ".join(line.split())[:300]
                break
        delivered_found = bool(status_field) or any(
            term in folded for term in DELIVERED_TERMS
        )
        if delivered_found and not status_field:
            status_field = next(
                (term for term in DELIVERED_TERMS if term in folded), ""
            )
        if not text.strip():
            details = "PDF没有可提取文本，需要人工查看或OCR"
        elif not tracking_found and not delivered_found:
            details = "未找到运单号和送达字段"
        elif not tracking_found:
            details = "未找到输入运单号或POD文件名中的主单号"
        elif not delivered_found:
            details = "未找到Delivered/Completed等送达字段"
    except Exception as exc:  # noqa: BLE001
        details = f"PDF读取失败：{exc}"

    if valid_pdf and tracking_found and delivered_found:
        result = "通过"
    elif valid_pdf:
        result = "人工复核"
    else:
        result = "失败"
    return PodAuditItem(
        tracking_number=str(tracking_number),
        carrier=str(carrier),
        pdf_file=str(path),
        sample_reason="",
        tracking_status="",
        valid_pdf=valid_pdf,
        tracking_found=tracking_found,
        delivered_found=delivered_found,
        status_field=status_field,
        result=result,
        details=details,
    )


def choose_pod_samples(
    results: Iterable[dict],
    sample_rates: Optional[dict] = None,
    rng: Optional[random.Random] = None,
) -> list[tuple[dict, str]]:
    configured = dict(DEFAULT_POD_AUDIT_RATES)
    if isinstance(sample_rates, dict):
        for carrier in configured:
            try:
                configured[carrier] = max(0, min(100, int(sample_rates.get(carrier, configured[carrier]))))
            except (TypeError, ValueError):
                pass
    eligible = {carrier: [] for carrier in configured}
    for result in results:
        pdf_file = str(result.get("POD文件") or result.get("pdf_file") or "").strip()
        carrier_text = str(
            result.get("快递公司") or result.get("carrier") or ""
        ).strip()
        carrier = next(
            (name for name in configured if name.casefold() == carrier_text.casefold()),
            "",
        )
        status = str(result.get("状态") or result.get("status") or "").strip().casefold()
        delivered = bool(result.get("is_delivered")) or status in {
            "delivered", "completed"
        }
        if not delivered or not pdf_file:
            continue
        if not carrier:
            continue
        if carrier == "FedEx":
            detail_file = str(
                result.get("POD详情文件") or result.get("detail_pdf_file") or ""
            ).strip()
            if (pdf_file and Path(pdf_file).is_file()) or (
                detail_file and Path(detail_file).is_file()
            ):
                eligible[carrier].append(result)
        elif Path(pdf_file).is_file():
            eligible[carrier].append(result)

    generator = rng or random.SystemRandom()
    selected = []
    for carrier, carrier_items in eligible.items():
        rate_percent = configured[carrier]
        if not carrier_items or rate_percent <= 0:
            continue
        count = min(
            math.ceil(len(carrier_items) * rate_percent / 100),
            len(carrier_items),
        )
        for source in generator.sample(carrier_items, count):
            if carrier == "FedEx":
                paths = (("POD文件", "查询主页"), ("POD详情文件", "详情页"))
            else:
                paths = (("POD文件", "POD"),)
            for key, label in paths:
                item = dict(source)
                item["_audit_pdf_file"] = str(item.get(key) or "")
                item["_audit_pod_kind"] = label
                item["_audit_sample_rate"] = rate_percent / 100
                selected.append((item, f"{carrier} {rate_percent}%随机抽查"))
    return selected


def _save_audit_workbook(items: list[PodAuditItem], output_file: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "POD抽查"
    sheet.sheet_view.showGridLines = False
    headers = (
        "运单号", "承运商", "POD类型", "抽查比例", "查询状态", "POD文件",
        "PDF有效", "运单号匹配", "PDF提取状态", "送达状态匹配", "结果", "说明",
    )
    sheet.append(headers)
    for item in items:
        sheet.append((
            item.tracking_number,
            item.carrier,
            item.pod_kind,
            item.sample_rate,
            item.tracking_status,
            item.pdf_file,
            "是" if item.valid_pdf else "否",
            "是" if item.tracking_found else "否",
            item.status_field,
            "是" if item.delivered_found else "否",
            item.result,
            item.details,
        ))
        sheet.cell(sheet.max_row, 4).number_format = "0%"
    header_fill = PatternFill("solid", fgColor="173F5F")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = (18, 12, 12, 12, 24, 42, 10, 12, 36, 14, 12, 44)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)


def audit_pod_sample(
    results: Iterable[dict],
    output_file,
    inspector: Callable[..., PodAuditItem] = inspect_pod,
    rng: Optional[random.Random] = None,
    sample_rates: Optional[dict] = None,
) -> list[PodAuditItem]:
    sampled = choose_pod_samples(results, sample_rates=sample_rates, rng=rng)
    items = []
    for result, reason in sampled:
        pdf_to_check = (
            result.get("_audit_pdf_file", "")
            if "_audit_pdf_file" in result
            else result.get("POD文件") or result.get("pdf_file")
        )
        checked = inspector(
            pdf_to_check,
            result.get("运单号") or result.get("tracking_number"),
            result.get("快递公司") or result.get("carrier", ""),
        )
        items.append(PodAuditItem(
            tracking_number=checked.tracking_number,
            carrier=checked.carrier,
            pdf_file=checked.pdf_file,
            sample_reason=reason,
            tracking_status=str(
                result.get("状态") or result.get("status") or ""
            ),
            valid_pdf=checked.valid_pdf,
            tracking_found=checked.tracking_found,
            delivered_found=checked.delivered_found,
            status_field=checked.status_field,
            result=checked.result,
            details=checked.details,
            pod_kind=str(result.get("_audit_pod_kind") or "POD"),
            sample_rate=float(result.get("_audit_sample_rate") or 0),
        ))
    if items:
        _save_audit_workbook(items, Path(output_file))
    return items
