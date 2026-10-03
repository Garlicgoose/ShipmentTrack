# -*- coding: utf-8 -*-
"""Persistent, atomic JSON queue for resumable FedEx webpage POD work."""
from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import datetime
import json
import os
from pathlib import Path
import threading
from typing import Iterable, Optional

from openpyxl import load_workbook

from modules.fedex_web_pod import is_valid_pdf, normalize_tracking_number, output_paths


READY_STATES = ("pending", "main_saved")
TERMINAL_STATES = ("completed", "skipped")
_LOCKS: dict[str, threading.RLock] = {}
_LOCKS_GUARD = threading.Lock()


@dataclass(frozen=True)
class FedExPodTask:
    number: str
    state: str
    main_pdf: str = ""
    detail_pdf: str = ""
    reason: str = ""
    source: str = ""
    position: int = 0
    updated_at: str = ""

    @property
    def expected_page(self) -> str:
        return "detail" if self.state == "main_saved" else "main"


def _now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _lock_for(path: Path) -> threading.RLock:
    key = str(path.resolve()).casefold()
    with _LOCKS_GUARD:
        return _LOCKS.setdefault(key, threading.RLock())


class FedExPodQueue:
    """Atomic JSON queue compatible with the dependency set shipped in v1.6."""

    def __init__(self, path) -> None:
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = _lock_for(self.path)
        with self._lock:
            if not self.path.exists():
                self._write({"version": 1, "next_position": 1, "tasks": []})

    def _read(self) -> dict:
        try:
            document = json.loads(self.path.read_text("utf-8"))
        except (OSError, ValueError) as exc:
            raise ValueError(f"FedEx POD 队列文件损坏：{self.path}") from exc
        if not isinstance(document, dict) or not isinstance(document.get("tasks"), list):
            raise ValueError(f"FedEx POD 队列格式无效：{self.path}")
        document.setdefault("version", 1)
        document.setdefault("next_position", len(document["tasks"]) + 1)
        return document

    def _write(self, document: dict) -> None:
        temporary = self.path.with_name(f".{self.path.name}.{os.getpid()}.tmp")
        try:
            temporary.write_text(
                json.dumps(document, ensure_ascii=False, indent=2), "utf-8"
            )
            temporary.replace(self.path)
        finally:
            temporary.unlink(missing_ok=True)

    @staticmethod
    def _to_task(item: dict) -> FedExPodTask:
        allowed = FedExPodTask.__dataclass_fields__
        values = {key: item.get(key, "") for key in allowed}
        values["position"] = int(values.get("position") or 0)
        return FedExPodTask(**values)

    def add_numbers(self, numbers: Iterable[object], source: str = "") -> list[str]:
        added = []
        with self._lock:
            document = self._read()
            existing = {str(item.get("number")) for item in document["tasks"]}
            seen = set()
            for raw in numbers:
                try:
                    number = normalize_tracking_number(raw)
                except ValueError:
                    continue
                if number in seen or number in existing:
                    continue
                seen.add(number)
                existing.add(number)
                task = FedExPodTask(
                    number=number,
                    state="pending",
                    source=str(source or ""),
                    position=int(document["next_position"]),
                    updated_at=_now(),
                )
                document["tasks"].append(asdict(task))
                document["next_position"] = int(document["next_position"]) + 1
                added.append(number)
            if added:
                self._write(document)
        return added

    def import_excel(self, path) -> list[str]:
        source = Path(path)
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            headers = {str(cell.value or "").strip(): cell.column for cell in sheet[1]}
            column = headers.get("运单号")
            if not column:
                raise ValueError("Excel 缺少“运单号”列表头")
            numbers = [
                sheet.cell(row, column).value
                for row in range(2, sheet.max_row + 1)
                if sheet.cell(row, column).value not in (None, "")
            ]
        finally:
            workbook.close()
        return self.add_numbers(numbers, source=str(source.resolve()))

    def tasks(self, states: Optional[Iterable[str]] = None) -> list[FedExPodTask]:
        wanted = set(states or ())
        with self._lock:
            tasks = [self._to_task(item) for item in self._read()["tasks"]]
        if wanted:
            tasks = [task for task in tasks if task.state in wanted]
        return sorted(tasks, key=lambda task: task.position)

    def get(self, number: str) -> Optional[FedExPodTask]:
        number = normalize_tracking_number(number)
        return next((task for task in self.tasks() if task.number == number), None)

    def next_ready(self) -> Optional[FedExPodTask]:
        return next(iter(self.tasks(READY_STATES)), None)

    def mark_main_saved(self, number: str, path: str) -> None:
        self._update(number, "main_saved", main_pdf=str(path), reason="")

    def mark_completed(self, number: str, main_pdf: str, detail_pdf: str) -> None:
        self._update(number, "completed", main_pdf=str(main_pdf),
                     detail_pdf=str(detail_pdf), reason="")

    def pause(self, number: str, reason: str) -> None:
        self._update(number, "paused", reason=str(reason or "等待人工处理"))

    def skip(self, number: str, reason: str = "用户跳过") -> None:
        self._update(number, "skipped", reason=reason)

    def resume(self, number: str) -> None:
        task = self.get(number)
        if task is None:
            raise KeyError(number)
        state = "main_saved" if task.main_pdf and is_valid_pdf(Path(task.main_pdf)) else "pending"
        self._update(number, state, reason="")

    def sync_existing_pdfs(self, output_dir) -> None:
        root = Path(output_dir)
        for task in self.tasks():
            main_pdf, detail_pdf = output_paths(root, task.number)
            if is_valid_pdf(main_pdf) and is_valid_pdf(detail_pdf):
                self.mark_completed(task.number, str(main_pdf), str(detail_pdf))
            elif is_valid_pdf(main_pdf) and task.state != "completed":
                self.mark_main_saved(task.number, str(main_pdf))

    def counts(self) -> dict[str, int]:
        result = {key: 0 for key in (
            "total", "pending", "main_saved", "completed", "paused", "skipped"
        )}
        for task in self.tasks():
            result[task.state] = result.get(task.state, 0) + 1
            result["total"] += 1
        return result

    def _update(self, number: str, state: str, **values) -> None:
        number = normalize_tracking_number(number)
        with self._lock:
            document = self._read()
            for item in document["tasks"]:
                if item.get("number") != number:
                    continue
                item["state"] = state
                item["updated_at"] = _now()
                for key in ("main_pdf", "detail_pdf", "reason"):
                    if key in values:
                        item[key] = str(values[key] or "")
                self._write(document)
                return
        raise KeyError(number)
