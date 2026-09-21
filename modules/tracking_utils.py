# -*- coding: utf-8 -*-
"""快递批量查询：配置、数据清洗、浏览器会话（Shipment Track 版）。

- 支持 DHL / DSV / EI / UPS（Playwright 有头浏览器）+ FedEx（官方 API）
- 输出 Excel 新增「状态」列（读取到的原始状态，用于排查 bug）
  和「备注」列（错误信息 / FedEx 子单超过 40 需人工核查标记）
"""
from pathlib import Path
import importlib
import math
import re

from openpyxl import load_workbook
from modules.real_browser import RealBrowserController
from units import detect_browser_path, get_data_path

TRACKING_COMPANY_COL_INDEX = 0
TRACKING_NUMBER_COL_INDEX = 1

TRACKING_OUTPUT_COLUMNS = [
    "运单号", "快递公司", "状态", "抵达时间", "用时(秒)",
    "POD文件", "POD详情文件", "POD抽查", "备注",
]

# FedEx 关联运单上限：API 最多只返回 40 条（含主单自身）
FEDEX_RELATED_LIMIT = 40

TRACKING_CARRIER_CONFIG = {
    "DHL": {
        "module": "modules.dhl_module",
        "query_func": "query_dhl_one",
        "warmup_func": "warm_up_dhl",
        "install_func": None,
        "login_func": None,
        "viewport": {"width": 1366, "height": 900},
        "pdf_subdir": "DHL",
        "locale": "en-US",
    },
    "DSV": {
        "module": "modules.dsv_module",
        "query_func": "query_dsv_one",
        "warmup_func": "warm_up_dsv",
        "install_func": None,
        "login_func": None,
        "viewport": {"width": 1500, "height": 900},
        "pdf_subdir": "DSV",
        "locale": "en-US",
        "force_english": True,
    },
    "EI": {
        "module": "modules.ei_module",
        "query_func": "query_expeditors_one",
        "warmup_func": "warm_up_ei",
        "install_func": None,
        "login_func": "login_expo",
        "viewport": {"width": 1500, "height": 900},
        "pdf_subdir": "EI",
        "locale": "en-US",
        "force_english": True,
    },
    "UPS": {
        "module": "modules.ups_module",
        "query_func": "query_ups_one",
        "warmup_func": "warm_up_ups",
        "install_func": "install_ups_assistant_blocker",
        "login_func": None,
        "viewport": {"width": 1366, "height": 900},
        "pdf_subdir": "UPS",
        "locale": "en-US",
    },
    # FedEx 状态走官方 API；已送达网页 POD 由真实 Edge 另行打印。
    "FedEx": {
        "module": "modules.fedex_module",
        "query_func": "query_fedex_one",
        "warmup_func": None,
        "install_func": None,
        "login_func": None,
        "viewport": None,
        "pdf_subdir": "FedEx",
        "locale": "en-US",
        "api_based": True,
    },
}


def normalize_tracking_arrival_date(value):
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""

    match = re.search(r"(\d{4})/(\d{1,2})/(\d{1,2})", text)
    if match:
        return f"{match.group(1)}/{int(match.group(2))}/{int(match.group(3))}"

    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if match:
        return f"{match.group(1)}/{int(match.group(2))}/{int(match.group(3))}"

    month_map = {
        "jan": 1, "january": 1,
        "feb": 2, "february": 2,
        "mar": 3, "march": 3,
        "apr": 4, "april": 4,
        "may": 5,
        "jun": 6, "june": 6,
        "jul": 7, "july": 7,
        "aug": 8, "august": 8,
        "sep": 9, "sept": 9, "september": 9,
        "oct": 10, "october": 10,
        "nov": 11, "november": 11,
        "dec": 12, "december": 12,
    }

    match = re.search(r"(\d{1,2})-([A-Za-z]{3,9})-(\d{4})", text, re.IGNORECASE)
    if match:
        day = int(match.group(1))
        month_num = month_map.get(match.group(2).lower())
        year = match.group(3)
        if month_num:
            return f"{year}/{month_num}/{day}"

    match = re.search(r"([A-Za-z]{3,9})\s+(\d{1,2}),?\s+(\d{4})", text, re.IGNORECASE)
    if match:
        month_num = month_map.get(match.group(1).lower())
        day = int(match.group(2))
        year = match.group(3)
        if month_num:
            return f"{year}/{month_num}/{day}"

    return text


def normalize_tracking_number(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""

    text = str(value).strip()

    if text.endswith(".0"):
        text = text[:-2]

    text = text.replace(" ", "")
    text = text.replace("\u3000", "")

    return text


def normalize_tracking_company_name(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""

    text = str(value).strip().upper()
    text = text.replace(" ", "")
    text = text.replace("-", "")
    text = text.replace("_", "")

    mapping = {
        "DHL": "DHL",
        "DHLEXPRESS": "DHL",
        "UPS": "UPS",
        "UNITEDPARCELSERVICE": "UPS",
        "DSV": "DSV",
        "EI": "EI",
        "EXPO": "EI",
        "EXPEDITORS": "EI",
        "EXPEDITOR": "EI",
        "EXPEDITORSINTERNATIONAL": "EI",
        "FEDEX": "FedEx",
        "FEDERALEXPRESS": "FedEx",
    }

    if text in mapping:
        return mapping[text]

    if "DHL" in text:
        return "DHL"
    if "UPS" in text:
        return "UPS"
    if "DSV" in text:
        return "DSV"
    if "FEDEX" in text or "FEDERALEXPRESS" in text:
        return "FedEx"
    if "EXPEDITORS" in text or "EXPEDITOR" in text or text == "EI":
        return "EI"

    return text


def prepare_tracking_input_rows(input_file):
    """读取并清洗跟踪 Excel，仅保留前两列，返回按承运商排序的字典列表。"""
    workbook = load_workbook(input_file, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or len(header) < 2:
            raise ValueError("Excel 至少需要两列：第一列快递公司，第二列运单号。")

        cleaned = []
        for row in rows:
            company_original = row[TRACKING_COMPANY_COL_INDEX] if len(row) > 0 else None
            tracking_original = row[TRACKING_NUMBER_COL_INDEX] if len(row) > 1 else None
            company = normalize_tracking_company_name(company_original)
            tracking_number = normalize_tracking_number(tracking_original)
            if not company or not tracking_number:
                continue
            cleaned.append({
                "快递公司原始值": company_original,
                "运单号原始值": tracking_original,
                "快递公司": company,
                "运单号": tracking_number,
            })
    finally:
        workbook.close()

    unsupported = sorted(
        {row["快递公司"] for row in cleaned} - set(TRACKING_CARRIER_CONFIG)
    )

    if unsupported:
        raise ValueError(
            "存在不支持的快递公司名称："
            + ", ".join(unsupported)
            + "\n支持：DHL / DSV / EI / UPS / FedEx"
        )

    cleaned.sort(key=lambda row: (row["快递公司"], row["运单号"]))
    return cleaned


def minimize_browser_window(page):
    """通过 CDP 把 Chromium 窗口最小化。

    比 --start-minimized 可靠（部分 Chromium 版本忽略该启动参数）。
    窗口最小化后页面照常加载/执行，不阻塞查询。
    """
    cdp = None
    try:
        cdp = page.context.new_cdp_session(page)
        info = cdp.send("Browser.getWindowForTarget")
        window_id = info.get("windowId")
        if window_id:
            cdp.send("Browser.setWindowBounds", {
                "windowId": window_id,
                "bounds": {"windowState": "minimized"},
            })
            return True
    except Exception:
        return False
    finally:
        if cdp is not None:
            try:
                cdp.detach()
            except Exception:
                pass
    return False


class TrackingCarrierSession:
    """DHL/DSV/EI/UPS session attached to installed Edge or Chrome."""

    def __init__(
        self,
        playwright,
        carrier,
        output_dir,
        ei_login_enabled,
        ei_email,
        ei_password,
        chrome_path="",
        browser_type="edge",
        browser_path="",
        minimize_browser=True,
        log_func=None,
        save_pdf=True,
        delivery_statuses=None,
    ):
        self.playwright = playwright
        self.carrier = carrier
        self.output_dir = Path(output_dir)
        self.ei_login_enabled = ei_login_enabled
        self.ei_email = ei_email
        self.ei_password = ei_password
        self.chrome_path = chrome_path
        self.browser_type = str(browser_type or "edge").casefold()
        self.browser_path = browser_path or chrome_path
        self.minimize_browser = minimize_browser
        self.log = log_func or (lambda msg: None)
        self.save_pdf = save_pdf
        self.delivery_statuses = delivery_statuses or {}

        self.browser = None
        self.browser_controller = None
        self.context = None
        self.page = None
        self.module = None
        self.config = TRACKING_CARRIER_CONFIG[carrier]

    def start(self):
        module_name = self.config["module"]
        self.module = importlib.import_module(module_name)

        if self.carrier == "EI":
            if hasattr(self.module, "EXPO_EMAIL"):
                self.module.EXPO_EMAIL = self.ei_email
            if hasattr(self.module, "EXPO_PASSWORD"):
                self.module.EXPO_PASSWORD = self.ei_password
            if hasattr(self.module, "EI_LOGIN_ENABLED"):
                self.module.EI_LOGIN_ENABLED = self.ei_login_enabled

        if hasattr(self.module, "CUSTOM_DELIVERED_STATUSES"):
            setattr(
                self.module,
                "CUSTOM_DELIVERED_STATUSES",
                tuple(self.delivery_statuses.get(self.carrier, ())),
            )

        pdf_dir = self.output_dir / "pdf" / self.config["pdf_subdir"]
        pdf_dir.mkdir(parents=True, exist_ok=True)

        if hasattr(self.module, "PDF_DIR"):
            setattr(self.module, "PDF_DIR", str(pdf_dir))

        if self.browser_type not in {"edge", "chrome"}:
            self.browser_type = "edge"
        if not (self.browser_path and Path(self.browser_path).is_file()):
            self.browser_path = detect_browser_path(self.browser_type)
        browser_name = "Microsoft Edge" if self.browser_type == "edge" else "Google Chrome"
        if not self.browser_path:
            raise FileNotFoundError(f"没有检测到 {browser_name}，请先在设置中选择浏览器路径。")
        profile_dir = (
            get_data_path() / "browser_profiles" / self.browser_type / self.carrier.casefold()
        )
        self.browser_controller = RealBrowserController(
            self.playwright,
            self.browser_path,
            profile_dir,
            browser_name=browser_name,
            locale=self.config.get("locale", "en-US"),
            minimize=self.minimize_browser,
            log_func=self.log,
        )
        self.context, self.page = self.browser_controller.start()
        self.browser = self.browser_controller.browser

        install_func_name = self.config.get("install_func")
        if install_func_name and hasattr(self.module, install_func_name):
            install_func = getattr(self.module, install_func_name)
            install_func(self.context)

        try:
            self.page.set_viewport_size(
                self.config.get("viewport", {"width": 1366, "height": 900})
            )
        except Exception:
            pass

        login_func_name = self.config.get("login_func")
        if login_func_name and hasattr(self.module, login_func_name):
            if self.carrier == "EI" and not self.ei_login_enabled:
                self.log("EI 未勾选登录，跳过登录步骤")
            else:
                self.log(f"{self.carrier} 执行登录函数：{login_func_name}")
                login_func = getattr(self.module, login_func_name)
                login_func(self.page)

        warmup_func_name = self.config.get("warmup_func")
        if warmup_func_name and hasattr(self.module, warmup_func_name):
            self.log(f"{self.carrier} 执行预热函数：{warmup_func_name}")
            warmup_func = getattr(self.module, warmup_func_name)
            warmup_func(self.page)

        # DSV / EI 官网没有把语言写进 URL，落地后如果仍是中文再兜底切英文。
        if self.config.get("force_english"):
            try:
                self.browser_controller.ensure_english_page(self.page)
            except Exception:
                pass

        # 预热后再统一清理一次残留弹窗/浮层，确保查询不被遮挡
        self._cleanup_overlays()
        self.log(f"{self.carrier} 预热完成，开始查询")

    def _cleanup_overlays(self):
        """快速清理弹窗/浮层（只调模块的 remove_overlays 兜底）。

        各模块 warmup/login 里已处理 cookie/welcome 弹窗；
        这里只做快速清理，避免重复调用慢速的 cookie 点击函数。
        """
        if not self.page or not self.module:
            return
        fn = getattr(self.module, "remove_overlays", None)
        if callable(fn):
            try:
                fn(self.page)
            except Exception:
                pass

    def query_one(self, tracking_number):
        # 查询前先清理一次可能新出现的弹窗
        self._cleanup_overlays()

        query_func_name = self.config["query_func"]

        if not hasattr(self.module, query_func_name):
            raise AttributeError(f"{self.config['module']} 中找不到函数：{query_func_name}")

        query_func = getattr(self.module, query_func_name)

        # 所有浏览器承运商使用统一签名。不要捕获 TypeError 后重新查询，
        # 否则模块内部真正的 TypeError 会导致同一运单被重复提交。
        return query_func(
            self.page,
            tracking_number,
            save_pdf=self.save_pdf,
        )

    def close(self):
        self.log(f"关闭 {self.carrier} 浏览器环境")

        if self.browser_controller is not None:
            self.browser_controller.close()

        self.page = None
        self.context = None
        self.browser = None
        self.browser_controller = None
        self.module = None
