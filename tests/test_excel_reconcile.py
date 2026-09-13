import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from modules.excel_reconcile import extract_date_from_name, merge_and_reconcile_excel
from modules.settings_store import FilenameMappingRule


def create_inspect(path: Path, quantities):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B", "C", "D", "E", "Packed quantity", "Double"])
    sheet.column_dimensions["A"].width = 24
    for index, quantity in enumerate(quantities, 1):
        row = sheet.max_row + 1
        sheet.append([f"row-{index}", "merged", "", "", "", quantity, f"=F{row}*2"])
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        sheet.row_dimensions[row].height = 28
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor="DDEBF7")
    workbook.save(path)


def create_droplist(path: Path, quantities):
    workbook = Workbook()
    first = workbook.active
    first.title = "Cover"
    data = workbook.create_sheet("Data")
    workbook.create_sheet("Summary")
    data.append([])
    data.append([])
    data.append(["S/O", "B", "C", "D", "E", "QTY"])
    for index, quantity in enumerate(quantities, 1):
        data.append([f"row-{index}", "", "", "", "", quantity])
    data.append(["Total", "", "", "", "", sum(quantities)])
    workbook.save(path)


def create_droplist_with_ignored_sheet1(path: Path, quantities, ignored_quantities):
    workbook = Workbook()
    workbook.active.title = "9.10"
    data = workbook.create_sheet("无类型和日期")
    ignored = workbook.create_sheet("Sheet1")
    workbook.create_sheet("Address")
    for sheet, values in ((data, quantities), (ignored, ignored_quantities)):
        sheet.append([])
        sheet.append([])
        sheet.append(["S/O", "B", "C", "D", "E", "QTY"])
        for index, quantity in enumerate(values, 1):
            sheet.append([f"row-{index}", "", "", "", "", quantity])
        sheet.append(["Total", "", "", "", "", sum(values)])
    workbook.save(path)


class ExcelReconcileTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.inspect = self.root / "inspect"
        self.droplist = self.root / "droplist"
        self.inspect.mkdir()
        self.droplist.mkdir()
        self.rules = [
            FilenameMappingRule("澳车", "光联", note="光联业务", display_type="澳车"),
            FilenameMappingRule("814S", "MPO", display_type="814S"),
            FilenameMappingRule("光联", "光联", display_type="光联"),
            FilenameMappingRule("MPO", "MPO", display_type="MPO"),
        ]

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_date_parser_supports_year_and_month_day(self):
        self.assertEqual(((0, 9, 10), "9.10"), extract_date_from_name("9.10光联.xlsx"))
        self.assertEqual(
            ((2026, 9, 10), "2026.9.10"),
            extract_date_from_name("2026-09-10 MPO.xlsx"),
        )
        self.assertEqual(((0, 0, 0), ""), extract_date_from_name("unknown.xlsx"))
        self.assertEqual(((0, 9, 10), "9.10"), extract_date_from_name("9.10"))

    def test_merge_and_reconcile_outputs_traceable_sheets(self):
        create_inspect(self.inspect / "9.10 澳车.xlsx", [10, 20])
        create_inspect(self.inspect / "9.10 814S.xlsx", [5])
        day = self.droplist / "9.10"
        day.mkdir()
        create_droplist(day / "Drop shipment list9.10光联.xlsx", [30])
        create_droplist(day / "Drop shipment list9.10MPO.xlsx", [4])
        output = self.root / "output"
        progress = []

        result = merge_and_reconcile_excel(
            self.inspect,
            self.droplist,
            output,
            self.rules,
            progress.append,
        )

        self.assertEqual(2, result.inspect_files)
        self.assertEqual(2, result.droplist_files)
        self.assertEqual(3, result.inspect_rows)
        self.assertEqual(2, result.droplist_rows)
        by_type = {row.target_type: row for row in result.rows}
        self.assertEqual("一致", by_type["光联"].result)
        self.assertEqual(0, by_type["光联"].difference)
        self.assertEqual("不一致", by_type["MPO"].result)
        self.assertEqual(1, by_type["MPO"].difference)
        self.assertEqual([45, 80, 100], progress)

        self.assertEqual(output / "合并检验表.xlsx", result.inspect_output_file)
        self.assertEqual(output / "合并Droplist.xlsx", result.droplist_output_file)
        self.assertTrue(result.droplist_output_file.is_file())
        workbook = load_workbook(result.inspect_output_file, data_only=False)
        self.assertEqual(
            ["合并检验表", "类型箱数", "核对汇总", "异常文件"],
            workbook.sheetnames,
        )
        self.assertEqual("类型", workbook["合并检验表"][1][7].value)
        self.assertEqual("归总类别", workbook["合并检验表"][1][8].value)
        inspect_rows = list(workbook["合并检验表"].iter_rows(min_row=2, values_only=True))
        guanglian = next(row for row in inspect_rows if row[7] == "澳车")
        self.assertEqual("光联", guanglian[8])
        self.assertEqual("光联业务", guanglian[11])
        self.assertFalse(workbook["核对汇总"].sheet_view.showGridLines)
        type_rows = list(
            workbook["类型箱数"].iter_rows(min_row=2, values_only=True)
        )
        self.assertIn(("9.10", "澳车", 30, "光联"), type_rows)
        self.assertIn(("9.10", "814S", 5, "MPO"), type_rows)
        self.assertFalse(workbook["类型箱数"].sheet_view.showGridLines)

        merged_sheet = workbook["合并检验表"]
        self.assertEqual(24, merged_sheet.column_dimensions["A"].width)
        self.assertEqual(28, merged_sheet.row_dimensions[2].height)
        self.assertEqual("00DDEBF7", merged_sheet["A2"].fill.fgColor.rgb)
        self.assertIn("B2:C2", {str(item) for item in merged_sheet.merged_cells.ranges})
        # 第二个文件的数据被追加到后续行，公式引用必须同步平移。
        self.assertEqual("=F3*2", merged_sheet["G3"].value)

    def test_unmatched_files_are_kept_and_reported(self):
        create_inspect(self.inspect / "9.11 未知客户.xlsx", [7])
        create_droplist(
            self.droplist / "Drop shipment list9.11MPO.xlsx",
            [7],
        )
        output = self.inspect
        result = merge_and_reconcile_excel(
            self.inspect,
            self.droplist,
            output,
            self.rules,
        )
        self.assertTrue(any(issue[0] == "检验表" for issue in result.issues))
        self.assertTrue(any(row.target_type == "未识别" for row in result.rows))

        # 第二次运行时不得把第一次输出再次当成输入。
        second = merge_and_reconcile_excel(
            self.inspect,
            self.droplist,
            output,
            self.rules,
        )
        self.assertEqual(1, second.inspect_files)

    def test_overseas_numbered_truck_defaults_to_guanglian(self):
        source = self.inspect / "9.12 国外出货第6车.xlsx"
        create_inspect(source, [8])
        workbook = load_workbook(source)
        workbook.active.cell(20, 1).fill = PatternFill("solid", fgColor="FFFFFF")
        workbook.save(source)
        create_droplist(
            self.droplist / "Drop shipment list9.12光联.xlsx",
            [8],
        )

        result = merge_and_reconcile_excel(
            self.inspect,
            self.droplist,
            self.root / "output",
            self.rules,
        )

        row = next(item for item in result.rows if item.target_type == "光联")
        self.assertEqual("一致", row.result)
        self.assertFalse(any(issue[1] == source.name for issue in result.issues))
        merged = load_workbook(result.inspect_output_file)["合并检验表"]
        self.assertEqual("光联", merged.cell(2, 8).value)
        self.assertEqual("光联", merged.cell(2, 9).value)
        self.assertEqual(2, merged.max_row)

    def test_droplist_uses_headers_and_ignores_sheet1_and_empty_files(self):
        create_inspect(self.inspect / "9.10 MPO国外EI自提.xlsx", [12])
        day = self.droplist / "9.10"
        day.mkdir()
        create_droplist_with_ignored_sheet1(
            day / "Drop shipment list without date MPO.xlsx",
            [12],
            [999],
        )
        Workbook().save(day / "Drop shipment list empty MPO.xlsx")

        result = merge_and_reconcile_excel(
            self.inspect,
            self.droplist,
            self.root / "output",
            self.rules,
        )

        self.assertEqual(1, result.droplist_files)
        self.assertEqual(1, result.droplist_rows)
        mpo = next(
            row for row in result.rows
            if row.target_type == "MPO" and row.date == "9.10"
        )
        self.assertEqual("9.10", mpo.date)
        self.assertEqual(12, mpo.droplist_quantity)
        self.assertEqual("一致", mpo.result)
        self.assertTrue(any("empty" in issue[1] for issue in result.issues))


if __name__ == "__main__":
    unittest.main()
