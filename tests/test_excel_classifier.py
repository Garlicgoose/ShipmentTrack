import unittest

from openpyxl import Workbook

from modules.excel_classifier import (
    classify_business_type,
    sample_identifiers,
    structure_fingerprint,
)


class ExcelClassifierTests(unittest.TestCase):
    def test_structure_has_priority_over_filename_hint(self):
        decision = classify_business_type(
            structural_category="MPO",
            samples=(),
            references={},
            filename_category="光联",
        )
        self.assertTrue(decision.confirmed)
        self.assertEqual("MPO", decision.category)
        self.assertEqual("工作表结构", decision.source)

    def test_reference_sample_overrides_conflicting_filename(self):
        decision = classify_business_type(
            samples=("AWB-002",),
            references={"光联": {"awb001"}, "MPO": {"awb002"}},
            filename_category="光联",
        )
        self.assertTrue(decision.confirmed)
        self.assertEqual("MPO", decision.category)

    def test_unique_reference_match_confirms_category(self):
        decision = classify_business_type(
            samples=("AWB-002", "unknown"),
            references={"光联": {"awb001"}, "MPO": {"awb002"}},
            filename_category="MPO",
        )
        self.assertEqual("MPO", decision.category)
        self.assertGreaterEqual(decision.confidence, 0.9)

    def test_sampler_stops_at_small_limit(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(("Tracking number", "QTY"))
        for index in range(100):
            sheet.append((f"AWB-{index:03}", index))
        values = sample_identifiers(sheet, limit=5)
        self.assertEqual(5, len(values))
        self.assertEqual("awb000", values[0])

    def test_structure_fingerprint_uses_content_not_filename(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "Business Unit"
        sheet["B1"] = "MPO"
        self.assertEqual("MPO", structure_fingerprint(sheet))

    def test_mixed_shipments_never_use_majority_or_filename_to_force_category(self):
        decision = classify_business_type(
            samples=("111111", "222222", "333333"),
            references={"光联": {"111111", "222222"}, "MPO": {"333333"}},
            filename_category="光联",
        )
        self.assertTrue(decision.conflict)
        self.assertFalse(decision.confirmed)

    def test_sales_orders_are_not_tracking_numbers(self):
        sheet = Workbook().active
        sheet.append(("S/O", "QTY"))
        sheet.append(("SHARED-SO", 1))
        self.assertEqual((), sample_identifiers(sheet))


if __name__ == "__main__":
    unittest.main()
