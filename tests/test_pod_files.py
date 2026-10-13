import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from modules.pod_files import move_pod_files, update_tracking_pod_paths


class PodFileTests(unittest.TestCase):
    def test_moves_into_target_root_and_skips_conflicts_without_approval(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            first_dir = root / "input-a"
            second_dir = root / "input-b"
            first_dir.mkdir()
            second_dir.mkdir()
            first = first_dir / "123.pdf"
            second = second_dir / "123.pdf"
            first.write_bytes(b"first")
            second.write_bytes(b"second")
            result = move_pod_files(
                [("DHL", str(first)), ("FedEx", str(second))], root / "archive"
            )
            targets = [Path(new) for _old, new in result.moved]
            self.assertEqual(1, len(targets))
            self.assertEqual({"123.pdf"}, {path.name for path in targets})
            self.assertTrue(all(path.parent == root / "archive" for path in targets))
            self.assertEqual({b"first"}, {path.read_bytes() for path in targets})
            self.assertTrue(first_dir.is_dir())
            self.assertTrue(second_dir.is_dir())
            self.assertFalse((root / "archive" / "DHL").exists())
            self.assertFalse((root / "archive" / "FedEx").exists())
            self.assertFalse(first.exists())
            self.assertEqual(b"second", second.read_bytes())
            self.assertEqual((str(second.resolve()),), result.skipped)

    def test_updates_both_pod_columns_in_tracking_result(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "tracking_result.xlsx"
            book = Workbook()
            book.active.append(("运单号", "POD文件", "POD详情文件"))
            book.active.append(("123", "C:/old/123.pdf", "C:/old/123+.pdf"))
            book.save(path)
            count = update_tracking_pod_paths(path, (
                ("C:/old/123.pdf", "D:/archive/123.pdf"),
                ("C:/old/123+.pdf", "D:/archive/123+.pdf"),
            ))
            self.assertEqual(2, count)
            output = load_workbook(path, read_only=True)
            self.assertEqual("D:/archive/123.pdf", output.active["B2"].value)
            self.assertEqual("D:/archive/123+.pdf", output.active["C2"].value)
            output.close()

    def test_never_moves_directories_or_non_pdf_source_files(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "source_A" / "tracking_1"
            source.mkdir(parents=True)
            pod = source / "POD.pdf"
            invoice = source / "invoice.xlsx"
            fake_pdf_directory = source / "folder.pdf"
            pod.write_bytes(b"%PDF-pod")
            invoice.write_bytes(b"invoice")
            fake_pdf_directory.mkdir()

            result = move_pod_files(
                [
                    ("FedEx", str(source)),
                    ("FedEx", str(fake_pdf_directory)),
                    ("FedEx", str(invoice)),
                    ("FedEx", str(pod)),
                ],
                root / "archive",
            )

            self.assertEqual(1, len(result.moved))
            self.assertEqual(3, len(result.skipped))
            self.assertTrue(source.is_dir())
            self.assertTrue(fake_pdf_directory.is_dir())
            self.assertTrue(invoice.is_file())
            self.assertFalse(pod.exists())
            self.assertTrue((root / "archive" / "POD.pdf").is_file())

    def test_main_and_plus_detail_pdfs_move_without_changing_existing_target(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "pdf" / "FedEx"
            target = root / "archive"
            source.mkdir(parents=True)
            target.mkdir()
            (target / "123.pdf").write_bytes(b"existing")
            (source / "123.pdf").write_bytes(b"main")
            (source / "123+.pdf").write_bytes(b"detail")
            result = move_pod_files([
                ("FedEx", str(source / "123.pdf")),
                ("FedEx", str(source / "123+.pdf")),
                ("FedEx", str(source / "123+.pdf")),
            ], target)
            self.assertEqual(1, len(result.moved))
            self.assertEqual(b"existing", (target / "123.pdf").read_bytes())
            self.assertEqual(b"main", (source / "123.pdf").read_bytes())
            self.assertFalse((target / "123 (1).pdf").exists())
            self.assertEqual(b"detail", (target / "123+.pdf").read_bytes())
            self.assertEqual([source / "123.pdf"], list(source.iterdir()))
            self.assertTrue(source.is_dir())
            again = move_pod_files([("FedEx", str(target / "123+.pdf"))], target)
            self.assertEqual((), again.moved)
            self.assertEqual(1, len(again.skipped))
            self.assertFalse((target / "123+ (1).pdf").exists())

    def test_all_other_moves_finish_before_one_overwrite_decision(self):
        for approve in (False, True):
            with self.subTest(approve=approve), tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                source, target = root / "FedEx", root / "archive"
                source.mkdir()
                target.mkdir()
                for name in ("123.pdf", "123+.pdf"):
                    (source / name).write_bytes(b"new")
                    (target / name).write_bytes(b"old")
                (source / "789.pdf").write_bytes(b"other")
                def decide(conflicts):
                    self.assertEqual(b"other", (target / "789.pdf").read_bytes())
                    self.assertFalse((source / "789.pdf").exists())
                    self.assertEqual(2, len(conflicts))
                    self.assertEqual(b"old", (target / "123.pdf").read_bytes())
                    return approve
                callback = mock.Mock(side_effect=decide)
                result = move_pod_files([
                    ("FedEx", str(source / name))
                    for name in ("123.pdf", "789.pdf", "123+.pdf")
                ], target, confirm_overwrite=callback)
                callback.assert_called_once()
                self.assertEqual(3 if approve else 1, len(result.moved))
                self.assertEqual(0 if approve else 2, len(result.skipped))
                self.assertEqual((), result.errors)
                for name in ("123.pdf", "123+.pdf"):
                    self.assertEqual(b"new" if approve else b"old", (target / name).read_bytes())
                    self.assertEqual(not approve, (source / name).exists())
                self.assertTrue(source.is_dir())

    def test_failed_overwrite_keeps_old_target_and_source(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source, target = root / "FedEx", root / "archive"
            source.mkdir()
            target.mkdir()
            (source / "123.pdf").write_bytes(b"new")
            (target / "123.pdf").write_bytes(b"old")
            with mock.patch("modules.pod_files.shutil.copy2", side_effect=OSError("copy failed")):
                result = move_pod_files(
                    [("FedEx", str(source / "123.pdf"))], target,
                    confirm_overwrite=lambda _conflicts: True,
                )
            self.assertEqual(1, len(result.errors))
            self.assertEqual(b"new", (source / "123.pdf").read_bytes())
            self.assertEqual(b"old", (target / "123.pdf").read_bytes())
            self.assertEqual([target / "123.pdf"], list(target.iterdir()))


if __name__ == "__main__":
    unittest.main()
