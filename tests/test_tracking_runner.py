import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from openpyxl import Workbook, load_workbook

from modules.tracking_runner import _aggregate_audit_results, run_tracking
from modules.tracking_utils import prepare_tracking_input_rows


class TrackingRunnerTests(unittest.TestCase):
    def test_fedex_two_page_audit_uses_worst_result(self):
        items = [
            SimpleNamespace(tracking_number="123", result="通过"),
            SimpleNamespace(tracking_number="123", result="失败"),
        ]
        self.assertEqual({"123": "失败"}, _aggregate_audit_results(items))

    def test_openpyxl_input_reader_normalizes_and_sorts(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(("快递公司", "运单号"))
            sheet.append(("UPS", " 1Z 123 "))
            sheet.append(("fedex", 123456789012))
            workbook.save(path)
            rows = prepare_tracking_input_rows(path)
        self.assertEqual(["FedEx", "UPS"], [row["快递公司"] for row in rows])
        self.assertEqual("123456789012", rows[0]["运单号"])
        self.assertEqual("1Z123", rows[1]["运单号"])

    def test_result_callback_receives_each_row_and_status_only_is_forwarded(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_file = root / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(("快递公司", "运单号"))
            sheet.append(("FedEx", "123456789012"))
            workbook.save(input_file)

            emitted = []
            progress = []
            fake_result = {
                "status": "Delivered",
                "is_delivered": "Y",
                "arrival_time": "2026-09-10T12:30:00+08:00",
                "error": "",
                "flag": "",
            }
            with mock.patch(
                "modules.tracking_runner.fedex_module.query_fedex_one",
                return_value=fake_result,
            ) as query, mock.patch(
                "modules.tracking_runner.fedex_module.close_shared_session"
            ) as close:
                output = run_tracking(
                    input_file=input_file,
                    output_dir=root / "output",
                    fedex_api_key="key",
                    fedex_api_secret="secret",
                    save_pdf=False,
                    progress=progress.append,
                    result=emitted.append,
                )

            self.assertEqual(1, len(emitted))
            self.assertEqual("Delivered", emitted[0]["状态"])
            self.assertEqual("2026/9/10", emitted[0]["抵达时间"])
            self.assertEqual([100], progress)
            self.assertFalse(query.call_args.kwargs["save_pdf"])
            close.assert_called_once()

            result_book = load_workbook(output, data_only=True)
            self.assertEqual("123456789012", str(result_book.active["A2"].value))

    def test_delivered_fedex_defers_web_pod_to_manual_panel(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_file = root / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(("快递公司", "运单号"))
            sheet.append(("FedEx", "519470439011"))
            workbook.save(input_file)
            emitted = []
            with mock.patch(
                "modules.tracking_runner.fedex_module.query_fedex_one",
                return_value={
                    "status": "Delivered", "is_delivered": "Y",
                    "arrival_time": "2026-09-09", "error": "", "flag": "",
                },
            ) as query, mock.patch(
                "playwright.sync_api.sync_playwright",
            ) as browser, mock.patch(
                "modules.tracking_runner.audit_pod_sample", return_value=[]
            ):
                output = run_tracking(
                    input_file=input_file,
                    output_dir=root / "output",
                    fedex_api_key="key",
                    fedex_api_secret="secret",
                    chrome_path=r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                    save_pdf=True,
                    result=emitted.append,
                )

            self.assertFalse(query.call_args.kwargs["save_pdf"])
            browser.assert_not_called()
            self.assertEqual("", emitted[0]["POD文件"])
            self.assertEqual("", emitted[0]["POD详情文件"])
            self.assertIn("待半自动保存", emitted[0]["备注"])
            result_book = load_workbook(output, data_only=True)
            headers = [cell.value for cell in result_book.active[1]]
            self.assertIn("POD详情文件", headers)

    def test_each_carrier_browser_closes_right_after_its_rows(self):
        """回归：每家公司查完必须立刻关掉浏览器，不要留窗口到桌面。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            input_file = root / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(("快递公司", "运单号"))
            sheet.append(("DHL", "1234567890"))
            sheet.append(("DHL", "1234567891"))
            sheet.append(("UPS", "1Z9999999999999999"))
            workbook.save(input_file)

            events = []

            def make_session(**kwargs):
                carrier = kwargs["carrier"]
                session = mock.Mock()
                session.query_one.return_value = {
                    "status": "In transit", "is_delivered": False,
                    "arrival_time": "", "error": "", "flag": "",
                }
                session.start.side_effect = lambda: events.append(f"start:{carrier}")
                session.close.side_effect = lambda: events.append(f"close:{carrier}")
                return session

            playwright_manager = mock.Mock()
            playwright_manager.start.return_value = mock.Mock()
            with mock.patch(
                "modules.tracking_runner.TrackingCarrierSession", side_effect=make_session
            ), mock.patch(
                "playwright.sync_api.sync_playwright", return_value=playwright_manager
            ), mock.patch(
                "modules.tracking_runner.fedex_module.close_shared_session"
            ):
                run_tracking(
                    input_file=input_file,
                    output_dir=root / "output",
                    save_pdf=False,
                    login_wait_seconds=0,
                )

        # 同一家公司两条单共用一次会话；两条查完立刻关闭，再开下家的浏览器
        self.assertEqual(
            ["start:DHL", "close:DHL", "start:UPS", "close:UPS"],
            events,
        )

    def test_company_login_wait_runs_once_for_first_web_carrier(self):
        """网页承运商第一次启动时等公司登录；FedEx API 不触发等待。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)

            def make_input():
                path = root / "input.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(("快递公司", "运单号"))
                sheet.append(("DHL", "1234567890"))
                sheet.append(("UPS", "1Z9999999999999999"))
                workbook.save(path)
                return path

            def make_session(**kwargs):
                session = mock.Mock()
                session.query_one.return_value = {
                    "status": "In transit", "is_delivered": False,
                    "arrival_time": "", "error": "", "flag": "",
                }
                return session

            playwright_manager = mock.Mock()
            playwright_manager.start.return_value = mock.Mock()

            # Edge：登录等待只出现一次
            edge_logs = []
            with mock.patch(
                "modules.tracking_runner.TrackingCarrierSession", side_effect=make_session
            ), mock.patch(
                "playwright.sync_api.sync_playwright", return_value=playwright_manager
            ), mock.patch("modules.tracking_runner.time.sleep") as sleep, mock.patch(
                "modules.tracking_runner.fedex_module.close_shared_session"
            ):
                run_tracking(
                    input_file=make_input(), output_dir=root / "out_edge",
                    save_pdf=False, browser_type="edge",
                    log=edge_logs.append, login_wait_seconds=7,
                )
            self.assertEqual(1, len([m for m in edge_logs if "公司要求登录" in m]))
            sleep.assert_any_call(7)

            # Chrome 也留同样的人工登录窗口；不是只给 Edge。
            chrome_logs = []
            with mock.patch(
                "modules.tracking_runner.TrackingCarrierSession", side_effect=make_session
            ), mock.patch(
                "playwright.sync_api.sync_playwright", return_value=playwright_manager
            ), mock.patch("modules.tracking_runner.time.sleep") as sleep, mock.patch(
                "modules.tracking_runner.fedex_module.close_shared_session"
            ):
                run_tracking(
                    input_file=make_input(), output_dir=root / "out_chrome",
                    save_pdf=False, browser_type="chrome",
                    log=chrome_logs.append, login_wait_seconds=7,
                )
            self.assertEqual(1, len([m for m in chrome_logs if "公司要求登录" in m]))
            sleep.assert_any_call(7)


if __name__ == "__main__":
    unittest.main()
