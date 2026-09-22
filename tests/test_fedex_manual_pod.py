import tempfile
import unittest
from pathlib import Path
from unittest import mock
from openpyxl import Workbook, load_workbook

from modules import fedex_manual_pod as manual


class ManualFedExPodTests(unittest.TestCase):
    def test_prepare_never_navigates_before_or_between_jobs(self):
        with tempfile.TemporaryDirectory() as temp_dir, \
             mock.patch.object(manual, "RealBrowserController") as controller_type:
            page = mock.Mock(url="about:blank")
            controller_type.return_value.start.return_value = (mock.Mock(pages=[page]), page)
            session = manual.ManualFedExPodSession(
                mock.Mock(), "msedge.exe", "edge", temp_dir
            )
            self.assertEqual("541964339019", session.prepare("5419 64339019"))
            self.assertFalse(session.fill_after_user_click("541964339019"))
            session.prepare("541964339020")
            page.click.assert_not_called()
            page.goto.assert_not_called()
            page.evaluate.assert_not_called()

    def test_user_chosen_fedex_locale_is_preserved_when_arming_input(self):
        with tempfile.TemporaryDirectory() as temp_dir, \
             mock.patch.object(manual, "RealBrowserController") as controller_type:
            page = mock.Mock(url="https://www.fedex.com/en-hk/tracking.html")
            controller_type.return_value.start.return_value = (mock.Mock(pages=[page]), page)
            session = manual.ManualFedExPodSession(mock.Mock(), "msedge.exe", "edge", temp_dir)
            session.prepare("541964339019")
            page.evaluate.side_effect = [None, ""]
            self.assertFalse(session.fill_after_user_click("541964339019"))
            script, number = page.evaluate.call_args_list[0].args
            self.assertEqual("541964339019", number)
            self.assertIn("pointerdown", script)
            self.assertIn("__shipmentTrackManualPending", script)
            self.assertNotIn("requestSubmit", script)
            page.goto.assert_not_called()

    def test_types_real_keys_only_after_user_click(self):
        with tempfile.TemporaryDirectory() as temp_dir, \
             mock.patch.object(manual, "RealBrowserController"):
            session = manual.ManualFedExPodSession(mock.Mock(), "msedge.exe", "edge", temp_dir)
            page = mock.Mock(url="https://www.fedex.com/en-hk/tracking.html")
            session.page = page
            session.context = mock.Mock(pages=[page])
            page.evaluate.side_effect = [None, "", "541964339019"]
            self.assertFalse(session.fill_after_user_click("541964339019"))
            page.keyboard.type.assert_not_called()
            self.assertTrue(session.fill_after_user_click("541964339019"))
            page.keyboard.press.assert_called_once_with("Control+A")
            page.keyboard.type.assert_called_once_with("541964339019", delay=60)

    def test_detail_requires_current_number_and_changed_page(self):
        with tempfile.TemporaryDirectory() as temp_dir, \
             mock.patch.object(manual, "RealBrowserController"):
            session = manual.ManualFedExPodSession(mock.Mock(), "msedge.exe", "edge", temp_dir)
            session.page = mock.Mock(url="https://www.fedex.com/fedextrack/")
            session.context = mock.Mock(pages=[session.page])
            before = ("541964339019\nDelivered", "https://www.fedex.com/fedextrack/")
            with mock.patch.object(manual, "_body_text", return_value="Travel History 541964339020" + "x" * 100):
                self.assertFalse(session.detail_ready("541964339019", before))
            with mock.patch.object(manual, "_body_text", return_value="Travel History 541964339019" + "x" * 100):
                self.assertTrue(session.detail_ready("541964339019", before))

    def test_blocked_page_never_prints(self):
        with tempfile.TemporaryDirectory() as temp_dir, \
             mock.patch.object(manual, "RealBrowserController"), \
             mock.patch.object(manual, "_print_current_page") as printer:
            session = manual.ManualFedExPodSession(mock.Mock(), "msedge.exe", "edge", temp_dir)
            session.page = mock.Mock(url="https://www.fedex.com/en-cn/tracking.html")
            session.context = mock.Mock(pages=[session.page])
            with mock.patch.object(manual, "_body_text", return_value="captcha 541964339019"):
                with self.assertRaisesRegex(RuntimeError, "验证码"):
                    session.save_main("541964339019")
            printer.assert_not_called()

    def test_user_navigation_during_listener_attach_is_transient(self):
        with tempfile.TemporaryDirectory() as temp_dir, \
             mock.patch.object(manual, "RealBrowserController"):
            session = manual.ManualFedExPodSession(mock.Mock(), "msedge.exe", "edge", temp_dir)
            page = mock.Mock(url="https://www.fedex.com/en-hk/tracking.html")
            session.context = mock.Mock(pages=[page])
            page.evaluate.side_effect = RuntimeError("Execution context was destroyed")
            self.assertFalse(session.fill_after_user_click("541964339019"))
            page.goto.assert_not_called()

    def test_manual_pod_paths_update_existing_tracking_excel(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "tracking_result.xlsx"
            book = Workbook()
            book.active.append(("运单号", "快递公司", "POD文件", "POD详情文件"))
            book.active.append(("541964339019", "FedEx", "", ""))
            book.save(path)
            self.assertTrue(manual.update_tracking_result_file(
                path, manual.ManualPodResult("541964339019", "main.pdf", "detail.pdf")
            ))
            output = load_workbook(path, read_only=True)
            self.assertEqual("main.pdf", output.active["C2"].value)
            self.assertEqual("detail.pdf", output.active["D2"].value)
            output.close()

    def test_refresh_audit_includes_manually_saved_fedex_pods(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            path = root / "tracking_result.xlsx"
            book = Workbook()
            book.active.append(("运单号", "快递公司", "状态", "POD文件", "POD详情文件", "POD抽查"))
            book.active.append(("541964339019", "FedEx", "Delivered", "main.pdf", "detail.pdf", ""))
            book.save(path)
            fake = mock.Mock(tracking_number="541964339019", result="通过")
            with mock.patch("modules.pod_audit.audit_pod_sample", return_value=[fake]) as audit:
                count = manual.refresh_pod_audit(path, root / "audit.xlsx", {"FedEx": 100})
            self.assertEqual(1, count)
            self.assertEqual(100, audit.call_args.kwargs["sample_rates"]["FedEx"])
            self.assertTrue(audit.call_args.args[0][0]["is_delivered"])
            output = load_workbook(path, read_only=True)
            self.assertEqual("通过", output.active["F2"].value)
            output.close()
